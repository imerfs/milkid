"""Reports page, Excel export and the financial-settings page (HTTP level)."""
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from main.models import CostCenter, Signatory
from main.tests.helpers import client_for, make_municipality
from main.tests.test_accounting_finance_panel import FinancePanelBase
from main.tests.test_accounting_panel import make_user


class ReportPageTests(FinancePanelBase):
    def test_every_report_renders(self):
        self.voucher()
        from main.accounting.reports import REPORTS
        for key in REPORTS:
            r = self.get('acc_panel', tab='reports', report=key)
            self.assertEqual(r.status_code, 200, key)
        self.assertContains(self.get('acc_panel', tab='reports', report='vouchers'), 'عوارض آزمون')
        self.assertContains(self.get('acc_panel', tab='reports', report='party_ledger'), 'طرف حساب را انتخاب کنید')

    def test_export_is_excel_and_tenant_scoped(self):
        self.voucher(1234)
        r = self.client.get(reverse('main:acc_p_report_export', args=[self.year.id, 'vouchers']))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], 'شماره')
        self.assertEqual(rows[1][6], 1234)
        other = make_user('dehyari')
        make_municipality(other, village='دیگر')
        r = client_for(other).get(reverse('main:acc_p_report_export', args=[self.year.id, 'vouchers']))
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.client.get(reverse('main:acc_p_report_export', args=[self.year.id, 'nope'])).status_code, 404)

    def test_export_bad_filter_is_400_text(self):
        r = self.client.get(reverse('main:acc_p_report_export', args=[self.year.id, 'vouchers']), {'date_from': 'x'})
        self.assertEqual(r.status_code, 400)

    def test_export_keeps_formula_like_text_inert(self):
        from main.accounting import ledger
        ledger.create_voucher(self.dehyar, self.year, self.d, 'general', '=HYPERLINK("http://x")', [
            {'account': self.bank, 'debit': 5, 'credit': 0}, {'account': self.tax, 'debit': 0, 'credit': 5}])
        r = self.client.get(reverse('main:acc_p_report_export', args=[self.year.id, 'vouchers']))
        ws = load_workbook(BytesIO(r.content)).active
        self.assertTrue(str(ws.cell(2, 5).value).startswith("'="))


class SettingsPageTests(FinancePanelBase):
    def test_add_and_delete_each_kind(self):
        self.assertContains(self.get('acc_panel', tab='settings'), 'مراکز هزینه')
        self.assertEqual(self.post('acc_p_setting_new', 'center', code='9', title='ویژه').status_code, 204)
        self.assertEqual(self.post('acc_p_setting_new', 'center', code='9', title='تکراری').status_code, 409)
        self.assertEqual(self.post('acc_p_setting_new', 'signatory', name='علی', position='دهیار').status_code, 204)
        self.assertEqual(self.post('acc_p_setting_new', 'description', kind='voucher', code='1', text='دریافت').status_code, 204)
        self.assertEqual(self.post('acc_p_setting_new', 'bogus').status_code, 404)
        center = CostCenter.objects.get(code='9')
        self.assertEqual(self.post('acc_p_setting_delete', 'center', center.id).status_code, 204)
        sig = Signatory.objects.get()
        other = make_user('dehyari')
        make_municipality(other, village='دیگر')
        self.assertEqual(self.post('acc_p_setting_delete', 'signatory', sig.id, client=client_for(other)).status_code, 404)
        self.assertEqual(self.post('acc_p_setting_delete', 'signatory', sig.id).status_code, 204)
