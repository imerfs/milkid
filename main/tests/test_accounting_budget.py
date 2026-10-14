"""Budget lines and budget-vs-actual: validation, closed years, tenant isolation, actuals from vouchers."""
from decimal import Decimal
from io import BytesIO

from django.db import IntegrityError, transaction
from django.urls import reverse
from openpyxl import load_workbook

from main.accounting import budget, ledger
from main.models import BudgetLine
from main.office.errors import Forbidden, Invalid, NotFound
from main.tests.helpers import client_for, make_municipality
from main.tests.test_accounting_ledger import LedgerBase, line, make_user


class BudgetRuleTests(LedgerBase):
    def test_set_and_update_budget(self):
        budget.set_budget(self.dehyar, self.year, self.tax, 5000)
        budget.set_budget(self.dehyar, self.year, self.tax, '7,000')
        self.assertEqual(BudgetLine.objects.get(fiscal_year=self.year, account=self.tax).amount, 7000)
        self.assertEqual(BudgetLine.objects.count(), 1)

    def test_only_budget_accounts(self):
        with self.assertRaises(Invalid):
            budget.set_budget(self.dehyar, self.year, self.bank, 10)

    def test_non_postable_group_account_refused(self):
        self.tax.is_postable = False
        self.tax.save()
        with self.assertRaises(Invalid):
            budget.set_budget(self.dehyar, self.year, self.tax, 10)

    def test_bad_amounts(self):
        for bad in (-1, '1.5', 'x', None, 10 ** 16):
            with self.assertRaises(Invalid, msg=repr(bad)):
                budget.set_budget(self.dehyar, self.year, self.tax, bad)

    def test_closed_year_refused(self):
        self.year.is_closed = True
        self.year.save()
        with self.assertRaises(Invalid):
            budget.set_budget(self.dehyar, self.year, self.tax, 10)

    def test_db_unique_per_year_and_account(self):
        BudgetLine.objects.create(fiscal_year=self.year, account=self.tax, amount=1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            BudgetLine.objects.create(fiscal_year=self.year, account=self.tax, amount=2)

    def test_other_municipality_cannot_touch(self):
        other = make_user('dehyari')
        make_municipality(other, village='روستای دیگر')
        with self.assertRaises(NotFound):
            budget.set_budget(other, self.year, self.tax, 10)
        with self.assertRaises(NotFound):
            budget.report(other, self.year)

    def test_citizen_forbidden(self):
        with self.assertRaises(Forbidden):
            budget.set_budget(make_user('citizen'), self.year, self.tax, 10)

    def test_year_must_belong_to_account_municipality(self):
        other = make_user('dehyari')
        make_municipality(other, village='روستای دیگر')
        other_year = ledger.open_year(other, 1404)
        with self.assertRaises(NotFound):
            budget.set_budget(self.dehyar, other_year, self.tax, 10)


class BudgetReportTests(LedgerBase):
    def test_income_actual_remaining_and_percent(self):
        budget.set_budget(self.dehyar, self.year, self.tax, 4000)
        self.voucher(1000)
        rows, totals = budget.report(self.dehyar, self.year)
        row = next(r for r in rows if r['account_id'] == self.tax.id)
        self.assertEqual((row['budget'], row['actual'], row['remaining']), (4000, 1000, 3000))
        self.assertEqual(row['percent'], 25)
        self.assertEqual(totals['resource']['budget'], 4000)
        self.assertEqual(totals['resource']['actual'], 1000)

    def test_expense_actual_is_debit_minus_credit(self):
        budget.set_budget(self.dehyar, self.year, self.salary, 900)
        ledger.create_voucher(self.dehyar, self.year, self.d, 'expense', 'حقوق',
                              [line(self.salary, debit=300), line(self.bank, credit=300)])
        rows, totals = budget.report(self.dehyar, self.year)
        row = next(r for r in rows if r['account_id'] == self.salary.id)
        self.assertEqual((row['actual'], row['remaining']), (300, 600))
        self.assertEqual(totals['use']['actual'], 300)

    def test_over_budget_is_flagged_not_blocked(self):
        budget.set_budget(self.dehyar, self.year, self.salary, 100)
        ledger.create_voucher(self.dehyar, self.year, self.d, 'expense', 'حقوق',
                              [line(self.salary, debit=300), line(self.bank, credit=300)])
        row = next(r for r in budget.report(self.dehyar, self.year)[0] if r['account_id'] == self.salary.id)
        self.assertTrue(row['over'])
        self.assertEqual(row['remaining'], -200)

    def test_actual_without_budget_still_listed_with_no_percent(self):
        self.voucher(1000)
        row = next(r for r in budget.report(self.dehyar, self.year)[0] if r['account_id'] == self.tax.id)
        self.assertEqual((row['budget'], row['actual'], row['percent']), (0, 1000, None))

    def test_final_only_ignores_drafts_and_deleted(self):
        budget.set_budget(self.dehyar, self.year, self.tax, 5000)
        draft = self.voucher(1000)
        done = self.voucher(200)
        ledger.finalize(self.dehyar, done)
        gone = self.voucher(50)
        ledger.delete_voucher(self.dehyar, gone)
        both = next(r for r in budget.report(self.dehyar, self.year)[0] if r['account_id'] == self.tax.id)
        final = next(r for r in budget.report(self.dehyar, self.year, include_temporary=False)[0]
                     if r['account_id'] == self.tax.id)
        self.assertEqual((both['actual'], final['actual']), (1200, 200))
        self.assertTrue(draft.pk)

    def test_other_tenants_vouchers_do_not_leak(self):
        other = make_user('dehyari')
        make_municipality(other, village='روستای دیگر')
        o_year = ledger.open_year(other, 1404)
        from main.accounting.chart import seed_default_accounts
        from main.models import Account
        seed_default_accounts(other)
        ob = Account.objects.get(municipality__admin_user=other, code='1120101')
        ot = Account.objects.get(municipality__admin_user=other, code='5010101001')
        ledger.create_voucher(other, o_year, self.d, 'income', 'x', [line(ob, debit=999), line(ot, credit=999)])
        rows, _ = budget.report(self.dehyar, self.year)
        self.assertEqual(rows, [])


class BudgetViewTests(LedgerBase):
    def setUp(self):
        super().setUp()
        self.client = client_for(self.dehyar)

    def test_page_and_save(self):
        url = reverse('main:acc_budget', args=[self.year.id])
        self.assertEqual(self.client.get(url).status_code, 200)
        r = self.client.post(url, {'account': self.tax.id, 'amount': '۴٬۰۰۰'})
        self.assertRedirects(r, url, fetch_redirect_response=False)
        self.assertEqual(BudgetLine.objects.get().amount, Decimal(4000))

    def test_bad_amount_is_400_and_saves_nothing(self):
        r = self.client.post(reverse('main:acc_budget', args=[self.year.id]), {'account': self.tax.id, 'amount': '1.5'})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(BudgetLine.objects.exists())

    def test_citizen_and_foreign_municipality(self):
        url = reverse('main:acc_budget', args=[self.year.id])
        self.assertEqual(client_for(make_user('citizen')).get(url).status_code, 302)
        other = make_user('dehyari')
        make_municipality(other, village='روستای دیگر')
        c = client_for(other)
        self.assertEqual(c.get(url).status_code, 404)
        self.assertEqual(c.post(url, {'account': self.tax.id, 'amount': '1'}).status_code, 404)
        self.assertEqual(c.get(reverse('main:acc_budget_export', args=[self.year.id])).status_code, 404)

    def test_excel_export(self):
        budget.set_budget(self.dehyar, self.year, self.tax, 4000)
        self.voucher(1000)
        r = self.client.get(reverse('main:acc_budget_export', args=[self.year.id]))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        ws = load_workbook(BytesIO(r.content)).active
        values = [[c.value for c in row] for row in ws.iter_rows()]
        self.assertTrue(any(4000 in row and 1000 in row for row in values))
