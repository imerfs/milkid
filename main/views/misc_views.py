from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, authenticate, logout
from django.contrib.auth.models import User
from django.contrib.auth.views import LoginView
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import UserCreationForm
from django.contrib import messages
from django.core.mail import send_mail
from django.conf import settings
from django.http import HttpResponseForbidden, JsonResponse, HttpResponse, FileResponse, Http404
from django.views.decorators.http import require_http_methods
from django.core.paginator import Paginator
from django.db.models import Q, Count, Sum
from django.db import transaction
from django.utils import timezone
from functools import wraps
import json
import os

from .common import role_required
from django.views.decorators.http import require_POST
from main.access import policy_required
from main.notifications import notify
from main.messaging_views import (
    messaging_dashboard, conversation_list, conversation_detail, 
    start_conversation, create_announcement, announcement_list, 
    announcement_detail, api_unread_count, api_send_quick_message
)
from main.forms import (ContactForm, CustomUserCreationForm, CitizenRegistrationForm, CustomAuthenticationForm, 
                   ServiceRequestForm, ProfileUpdateForm, DocumentUploadForm, ServiceWizardStepForm,
                   MunicipalityRegistrationForm, PropertyRegistrationForm, PropertySearchForm, ConstructionPermitForm,
                   PropertyAdvertisementForm)
from main.models import (UserProfile, ServiceRequest, Notification, RequestDocument, Service, ServiceStep, Property, Municipality,
                   PropertyAdvertisement, AutoCADFile, PropertyData, Payment, Conversation, Message, MessageRecipient,
                   ViolationReport, Inspection, InspectionDocument, TechnicalReport, Province, County, District, CityVillage,
                   Vote, CouncilComment, PropertyAdvertisementEditLog)


def about(request):
    """About page"""
    return render(request, 'main/about.html')


def contact(request):
    """Contact page"""
    if request.method == 'POST':
        form = ContactForm(request.POST)
        if form.is_valid():
            # Process contact form
            messages.success(request, 'پیام شما با موفقیت ارسال شد.')
            return redirect('main:contact')
    else:
        form = ContactForm()

    context = {'form': form}
    return render(request, 'main/contact.html', context)


def ai_fraud_detection(request):
    """AI-powered fraud detection"""
    return render(request, 'main/ai_fraud_detection.html')


def payments_debts(request):
    """Payments and debts management"""
    user_payments = Payment.objects.filter(user=request.user).order_by('-created_at')

    # Pagination
    from django.core.paginator import Paginator
    paginator = Paginator(user_payments, 10)
    page_number = request.GET.get('page')
    payments_page = paginator.get_page(page_number)

    context = {
        'payments': payments_page,
    }
    return render(request, 'main/payments_debts.html', context)


def process_payment(request, property_id):
    """Process payment for property"""
    property_obj = get_object_or_404(Property, id=property_id)

    # Create payment record
    payment = Payment.objects.create(
        user=request.user,
        property=property_obj,
        amount=property_obj.annual_tax,
        payment_type='property_tax',
        description=f'پرداخت مالیات سالانه برای پلاک {property_obj.parcel_number}',
        status='pending'
    )

    messages.success(request, 'درخواست پرداخت با موفقیت ثبت شد.')
    return redirect('main:payments_debts')


@policy_required('payment.review', api=True)
@require_POST
def approve_payment(request, payment_id):
    """Approve a payment"""
    payment = get_object_or_404(Payment, id=payment_id)
    payment.status = 'completed'
    payment.paid_date = timezone.now()
    payment.save()

    notify(payment.user, 'پرداخت تایید شد',
           f'پرداخت {payment.get_payment_type_display()} به مبلغ {payment.amount} تومان تایید شد.',
           kind='payment')

    return JsonResponse({'success': True, 'message': 'پرداخت تایید شد'})


@policy_required('payment.review', api=True)
@require_POST
def reject_payment(request, payment_id):
    """Reject a payment"""
    payment = get_object_or_404(Payment, id=payment_id)
    payment.status = 'rejected'
    payment.save()

    notify(payment.user, 'پرداخت رد شد',
           f'پرداخت {payment.get_payment_type_display()} به مبلغ {payment.amount} تومان رد شد.',
           kind='payment')

    return JsonResponse({'success': True, 'message': 'پرداخت رد شد'})


def announcements(request):
    """Public announcements"""
    return render(request, 'main/announcements.html')


def manage_announcements(request):
    """Manage announcements"""
    return render(request, 'main/manage_announcements.html')


def report_violation(request):
    """Report a violation"""
    return render(request, 'main/report_violation.html')


def my_violations(request):
    """User's violation reports"""
    return render(request, 'main/my_violations.html')


def suggestions(request):
    """Submit suggestions"""
    return render(request, 'main/suggestions.html')


def my_suggestions(request):
    """User's suggestions"""
    return render(request, 'main/my_suggestions.html')


def event_permits(request):
    """Event permits"""
    return render(request, 'main/event_permits.html')


def my_event_permits(request):
    """User's event permits"""
    return render(request, 'main/my_event_permits.html')


def local_businesses(request):
    """Local businesses directory"""
    from .models import LocalBusiness
    businesses = LocalBusiness.objects.filter(is_verified=True).order_by('-created_at')

    context = {
        'businesses': businesses,
    }
    return render(request, 'main/local_businesses.html', context)


def faq_center(request):
    """FAQ Center"""
    return render(request, 'main/faq.html')


def autocad_upload(request):
    """Upload AutoCAD file"""
    from .services.autocad_parser import AutoCADParser

    if request.method == 'POST':
        if 'autocad_file' in request.FILES:
            autocad_file = request.FILES['autocad_file']
            
            autocad_obj = AutoCADFile.objects.create(
                user=request.user,
                file=autocad_file,
                original_filename=autocad_file.name
            )

            try:
                parser = AutoCADParser(autocad_obj.file.path)
                property_data = parser.parse()

                PropertyData.objects.create(
                    autocad_file=autocad_obj,
                    data=property_data
                )

                autocad_obj.status = 'processed'
                autocad_obj.save()

                return redirect('main:autocad_result', pk=autocad_obj.pk)
            except Exception as e:
                autocad_obj.status = 'failed'
                autocad_obj.save()
                messages.error(request, f'خطا در پردازش فایل: {str(e)}')

    return render(request, 'main/autocad/upload.html')


def autocad_result(request, pk):
    """View AutoCAD parsing result"""
    autocad_file = get_object_or_404(AutoCADFile, pk=pk, user=request.user)

    context = {
        'autocad_file': autocad_file,
    }
    return render(request, 'main/autocad/result.html', context)


def autocad_history(request):
    """AutoCAD upload history"""
    autocad_files = AutoCADFile.objects.filter(user=request.user).order_by('-uploaded_at')

    # Pagination
    from django.core.paginator import Paginator
    paginator = Paginator(autocad_files, 10)
    page_number = request.GET.get('page')
    files_page = paginator.get_page(page_number)

    context = {
        'files': files_page,
    }
    return render(request, 'main/autocad/history.html', context)


def export_transactions_excel(request):
    """Export transactions to Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse
    import io

    try:
        date_from = request.GET.get('date_from', '')
        date_to = request.GET.get('date_to', '')
        section = request.GET.get('section', 'all')

        payments = Payment.objects.all()

        if date_from:
            try:
                from datetime import datetime
                from_date = datetime.strptime(date_from, '%Y-%m-%d').date()
                payments = payments.filter(created_at__date__gte=from_date)
            except:
                pass

        if date_to:
            try:
                from datetime import datetime
                to_date = datetime.strptime(date_to, '%Y-%m-%d').date()
                payments = payments.filter(created_at__date__lte=to_date)
            except:
                pass

        if section == 'income':
            payments = payments.filter(payment_type__in=['property_tax', 'permit_fee'])
        elif section == 'expense':
            payments = payments.filter(payment_type__in=['construction_fine', 'service_fee', 'other'])

        wb = Workbook()
        ws = wb.active
        ws.title = 'تراکنش‌ها'
        ws.right_to_left = True
        
        headers = ['ردیف', 'شماره رسید', 'نوع پرداخت', 'مبلغ (تومان)', 'شرح', 'وضعیت', 'تاریخ ایجاد', 'تاریخ پرداخت', 'نام کاربر']
        header_fill = PatternFill(start_color='0284c7', end_color='0284c7', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for idx, payment in enumerate(payments.order_by('-created_at'), start=2):
            ws.cell(row=idx, column=1, value=idx-1)
            ws.cell(row=idx, column=2, value=payment.receipt_number)
            ws.cell(row=idx, column=3, value=payment.get_payment_type_display())
            ws.cell(row=idx, column=4, value=int(payment.amount))
            ws.cell(row=idx, column=5, value=payment.description)
            ws.cell(row=idx, column=6, value=payment.get_status_display())
            ws.cell(row=idx, column=7, value=payment.created_at.strftime('%Y-%m-%d'))
            ws.cell(row=idx, column=8, value=payment.paid_date.strftime('%Y-%m-%d') if payment.paid_date else '-')
            ws.cell(row=idx, column=9, value=payment.user.get_full_name())
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=transactions.xlsx'
        return response

    except Exception as e:
        messages.error(request, f'خطا در ایجاد فایل اکسل: {str(e)}')
        return redirect('main:accountant_dashboard')


def api_notifications_list(request):
    """API endpoint to get user notifications"""
    try:
        limit = int(request.GET.get('limit', 10))
        unread_only = request.GET.get('unread_only', 'false').lower() == 'true'
        
        notifications = Notification.objects.filter(user=request.user)
        
        if unread_only:
            notifications = notifications.filter(is_read=False)
        
        notifications = notifications.order_by('-created_at')[:limit]
        
        data = [{
            'id': notif.id,
            'title': notif.title,
            'message': notif.message,
            'type': notif.notification_type,
            'is_read': notif.is_read,
            'created_at': notif.created_at.strftime('%Y-%m-%d %H:%M'),
        } for notif in notifications]
        
        return JsonResponse({
            'success': True,
            'notifications': data,
            'unread_count': Notification.objects.filter(user=request.user, is_read=False).count()
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_notification_mark_read(request, notification_id):
    """API endpoint to mark notification as read"""
    try:
        notification = get_object_or_404(Notification, id=notification_id, user=request.user)
        notification.is_read = True
        notification.save()
        
        return JsonResponse({
            'success': True,
            'message': 'اعلان به عنوان خوانده شده علامت‌گذاری شد'
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_notifications_mark_all_read(request):
    """API endpoint to mark all notifications as read"""
    try:
        Notification.objects.filter(user=request.user, is_read=False).update(is_read=True)
        
        return JsonResponse({
            'success': True,
            'message': 'همه اعلان‌ها به عنوان خوانده شده علامت‌گذاری شدند'
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_notification_delete(request, notification_id):
    """API endpoint to delete a notification"""
    try:
        notification = get_object_or_404(Notification, id=notification_id, user=request.user)
        notification.delete()
        
        return JsonResponse({
            'success': True,
            'message': 'اعلان حذف شد'
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_messages_list(request):
    """API endpoint to get user conversations"""
    try:
        conversations = Conversation.objects.filter(
            participants=request.user
        ).order_by('-updated_at')[:20]
        
        data = []
        for conv in conversations:
            last_message = conv.messages.order_by('-created_at').first()
            unread_count = MessageRecipient.objects.filter(
                message__conversation=conv,
                recipient=request.user,
                is_read=False
            ).count()
            
            data.append({
                'id': conv.id,
                'subject': conv.subject,
                'last_message': last_message.body[:50] if last_message else '',
                'last_sender': last_message.sender.get_full_name() if last_message else '',
                'updated_at': conv.updated_at.strftime('%Y-%m-%d %H:%M'),
                'unread_count': unread_count,
            })
        
        return JsonResponse({
            'success': True,
            'conversations': data
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_conversation_messages(request, conversation_id):
    """API endpoint to get messages in a conversation"""
    try:
        conversation = get_object_or_404(
            Conversation,
            id=conversation_id,
            participants=request.user
        )
        
        messages = conversation.messages.order_by('created_at')
        
        data = [{
            'id': msg.id,
            'body': msg.body,
            'sender': msg.sender.get_full_name(),
            'sender_id': msg.sender.id,
            'created_at': msg.created_at.strftime('%Y-%m-%d %H:%M'),
            'is_mine': msg.sender == request.user,
        } for msg in messages]
        
        MessageRecipient.objects.filter(
            message__conversation=conversation,
            recipient=request.user,
            is_read=False
        ).update(is_read=True, read_at=timezone.now())
        
        return JsonResponse({
            'success': True,
            'messages': data,
            'subject': conversation.subject
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_message_send(request):
    """API endpoint to send a new message"""
    import json
    try:
        data = json.loads(request.body)
        recipient_id = data.get('recipient_id')
        subject = data.get('subject', '').strip()
        body = data.get('body', '').strip()
        
        if not recipient_id or not body:
            return JsonResponse({
                'success': False,
                'message': 'گیرنده و متن پیام الزامی است'
            }, status=400)
        
        recipient = get_object_or_404(User, id=recipient_id)
        
        conversation = Conversation.objects.create(subject=subject or 'بدون موضوع')
        conversation.participants.add(request.user, recipient)
        
        message = Message.objects.create(
            conversation=conversation,
            sender=request.user,
            body=body
        )
        
        MessageRecipient.objects.create(
            message=message,
            recipient=recipient,
            is_read=False
        )
        
        Notification.objects.create(
            user=recipient,
            notification_type='new_message',
            title='پیام جدید',
            message=f'{request.user.get_full_name()}: {body[:50]}...'
        )
        
        return JsonResponse({
            'success': True,
            'message': 'پیام با موفقیت ارسال شد',
            'conversation_id': conversation.id
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_conversation_reply(request, conversation_id):
    """API endpoint to reply to an existing conversation"""
    import json
    try:
        conversation = get_object_or_404(
            Conversation, 
            id=conversation_id,
            participants=request.user
        )
        
        data = json.loads(request.body)
        body = data.get('body', '').strip()
        
        if not body:
            return JsonResponse({
                'success': False,
                'message': 'متن پیام الزامی است'
            }, status=400)
        
        message = Message.objects.create(
            conversation=conversation,
            sender=request.user,
            body=body
        )
        
        for participant in conversation.participants.exclude(id=request.user.id):
            MessageRecipient.objects.create(
                message=message,
                recipient=participant,
                is_read=False
            )
            
            Notification.objects.create(
                user=participant,
                notification_type='new_message',
                title=f'پاسخ به: {conversation.subject}',
                message=f'{request.user.get_full_name()}: {body[:50]}...'
            )
        
        conversation.save()
        
        return JsonResponse({
            'success': True,
            'message': 'پاسخ با موفقیت ارسال شد',
            'message_id': message.id
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_message_mark_read(request, message_id):
    """API endpoint to mark a message as read"""
    try:
        message_recipient = get_object_or_404(
            MessageRecipient,
            message_id=message_id,
            recipient=request.user
        )
        
        message_recipient.is_read = True
        message_recipient.read_at = timezone.now()
        message_recipient.save()
        
        return JsonResponse({
            'success': True,
            'message': 'پیام به عنوان خوانده شده علامت‌گذاری شد'
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def api_get_provinces(request):
    provinces = Province.objects.all().order_by('name').values('id', 'name')
    return JsonResponse(list(provinces), safe=False)

def api_get_counties(request, province_id=None):
    prov_id = province_id or request.GET.get('province_id')
    if prov_id:
        counties = County.objects.filter(province_id=prov_id).order_by('name').values('id', 'name')
    else:
        counties = County.objects.all().order_by('name').values('id', 'name')
    return JsonResponse(list(counties), safe=False)

def api_get_districts(request, county_id=None):
    c_id = county_id or request.GET.get('county_id')
    if c_id:
        districts = District.objects.filter(county_id=c_id).order_by('name').values('id', 'name')
    else:
        districts = District.objects.all().order_by('name').values('id', 'name')
    return JsonResponse(list(districts), safe=False)

def api_get_dehestans(request, district_id=None):
    from main.models import Dehestan
    d_id = district_id or request.GET.get('district_id')
    if d_id:
        dehestans = Dehestan.objects.filter(district_id=d_id).order_by('name').values('id', 'name')
    else:
        dehestans = Dehestan.objects.all().order_by('name').values('id', 'name')
    return JsonResponse(list(dehestans), safe=False)

def api_get_city_villages(request, district_id=None):
    d_id = district_id or request.GET.get('district_id')
    dehestan_id = request.GET.get('dehestan_id')
    query = CityVillage.objects.all()
    if d_id:
        query = query.filter(district_id=d_id)
    if dehestan_id:
        query = query.filter(dehestan_id=dehestan_id)
    settlements = query.order_by('name').values('id', 'name', 'settlement_type')
    return JsonResponse(list(settlements), safe=False)


def api_request_city_village(request):
    """ثبت درخواست افزودن شهر یا روستای جدید"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'روش درخواست نامعتبر است.'}, status=405)
    
    from main.models import Province, County, District, Dehestan, CityVillageRequest, CityVillage
    import json
    
    try:
        data = json.loads(request.body) if request.body else request.POST
        
        province_id = data.get('province_id')
        county_id = data.get('county_id')
        district_id = data.get('district_id')
        name = (data.get('name') or '').strip()
        settlement_type = data.get('settlement_type', 'village')
        requester_name = (data.get('requester_name') or '').strip()
        requester_phone = (data.get('requester_phone') or '').strip()
        description = (data.get('description') or '').strip()
        
        if not name:
            return JsonResponse({'success': False, 'message': 'نام شهر یا روستا الزامی است.'}, status=400)
            
        if not district_id:
            return JsonResponse({'success': False, 'message': 'انتخاب بخش الزامی است.'}, status=400)
            
        district = District.objects.get(id=district_id)
        county = district.county
        province = county.province
        
        # Check if already exists in CityVillage
        existing = CityVillage.objects.filter(district=district, name=name).first()
        if existing:
            return JsonResponse({
                'success': True,
                'already_exists': True,
                'item': {'id': existing.id, 'name': existing.name, 'settlement_type': existing.settlement_type},
                'message': 'این شهر/روستا در پایگاه داده موجود است و انتخاب گردید.'
            })
            
        # Create request or auto-add
        req = CityVillageRequest.objects.create(
            province=province,
            county=county,
            district=district,
            name=name,
            settlement_type=settlement_type,
            requester_name=requester_name,
            requester_phone=requester_phone,
            description=description,
            status='approved'  # Auto approve to immediately make available for selection
        )
        
        # Also create CityVillage so user can immediately select it
        cv = CityVillage.objects.create(
            district=district,
            name=name,
            settlement_type=settlement_type
        )
        
        return JsonResponse({
            'success': True,
            'item': {
                'id': cv.id,
                'name': cv.name,
                'settlement_type': cv.settlement_type
            },
            'message': f'درخواست با موفقیت ثبت شد و «{name}» به لیست شهرها و روستاهای این بخش اضافه شد.'
        })
    except District.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'بخش مورد نظر یافت نشد.'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'message': f'خطا: {str(e)}'}, status=500)


def get_counties(request):
    return api_get_counties(request)

def get_districts(request):
    return api_get_districts(request)

def get_settlements(request):
    return api_get_city_villages(request)

def test_geographic(request):
    return render(request, 'main/test_geographic.html')

def test_jalali_dates(request):
    return JsonResponse({'status': 'ok'})

